import fitz
from docx import Document as DocxDocument
from langchain_core.documents import Document
from pptx import Presentation

from rag_demo.loaders.docxloader import DocxLoader
from rag_demo.loaders.pdf_loader import PDFLoader
from rag_demo.loaders.pptx_loader import PptxLoader


def _make_pdf(path):
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Chapter 1\nThis document explains retrieval augmented generation and vector search.",
    )
    document.save(path)
    document.close()


def test_pdf_loader_extracts_chunks_and_metadata(tmp_path):
    path = tmp_path / "sample.pdf"
    _make_pdf(path)

    chunks = PDFLoader(str(path), chunk_size=100, chunk_overlap=10).load()

    assert chunks
    assert all(chunk.metadata["file_type"] == "pdf" for chunk in chunks)
    assert all(chunk.metadata["source"] == "sample.pdf" for chunk in chunks)
    assert all("chunk_index" in chunk.metadata for chunk in chunks)


def test_docx_loader_extracts_chunks_and_metadata(tmp_path):
    path = tmp_path / "sample.docx"
    document = DocxDocument()
    document.add_paragraph("Chapter 1")
    document.add_paragraph("Retrieval augmented generation combines search with language models. " * 12)
    document.save(path)

    chunks = DocxLoader(str(path), chunk_size=60, chunk_overlap=10).load()

    assert chunks
    assert all(chunk.metadata["file_type"] == "docx" for chunk in chunks)
    assert all(chunk.metadata["source"] == "sample.docx" for chunk in chunks)


def test_pptx_loader_uses_requested_chunk_size_and_metadata(tmp_path, monkeypatch):
    path = tmp_path / "sample.pptx"
    Presentation().save(path)
    loader = PptxLoader(str(path), chunk_size=40, chunk_overlap=5)
    monkeypatch.setattr(
        loader,
        "_extract_text",
        lambda: iter([Document(page_content="Retrieval augmented generation. " * 20, metadata={})]),
    )

    chunks = loader.load()

    assert len(chunks) > 1
    assert all(chunk.metadata["file_type"] == "pptx" for chunk in chunks)
    assert all(chunk.metadata["source"] == "sample.pptx" for chunk in chunks)
